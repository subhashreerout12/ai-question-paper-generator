"""Write a paper (and its answer key) to an editable Word file."""

import io
import os

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt


def _centered(doc: Document, text: str, size: int, bold: bool = False) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    run.bold = bold
    run.font.size = Pt(size)


def export_docx(
    paper: dict,
    path: str,
    title: str = "Question Paper",
    subject: str = "",
    duration: str = "1 Hour",
    include_answers: bool = False,
) -> None:
    """Save the paper. With include_answers=True the file becomes the answer key."""
    doc = Document()
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(11)

    _centered(doc, title + (" - ANSWER KEY" if include_answers else ""), 16, bold=True)
    if subject:
        _centered(doc, subject, 12)
    p = doc.add_paragraph(f"Time: {duration}\t\t\t\tMaximum Marks: {paper['total_marks']}")
    p.runs[0].bold = True

    if not include_answers:
        doc.add_paragraph("General instructions: Answer all questions. Marks are shown against each question.")

    number = 1

    if paper["mcq"]:
        marks = sum(q["marks"] for q in paper["mcq"])
        doc.add_heading(f"Section A - Multiple Choice Questions ({marks} marks)", level=2)
        for q in paper["mcq"]:
            doc.add_paragraph(f"Q{number}. {q['question']}  [{q['marks']}]").runs[0].bold = True
            for letter, option in zip("ABCD", q["options"]):
                opt = doc.add_paragraph(f"({letter}) {option}")
                opt.paragraph_format.left_indent = Pt(24)
                opt.paragraph_format.space_after = Pt(0)
            if include_answers:
                ans = doc.add_paragraph(f"Answer: {q['answer']}")
                ans.paragraph_format.left_indent = Pt(24)
                if q.get("explanation"):
                    ans.add_run(f"  -  {q['explanation']}").italic = True
            number += 1

    if paper["short"]:
        marks = sum(q["marks"] for q in paper["short"])
        doc.add_heading(f"Section B - Short Answer Questions ({marks} marks)", level=2)
        for q in paper["short"]:
            doc.add_paragraph(f"Q{number}. {q['question']}  [{q['marks']}]").runs[0].bold = True
            if include_answers:
                a = doc.add_paragraph(f"Model answer: {q['answer']}")
                a.paragraph_format.left_indent = Pt(24)
                if q.get("key_points"):
                    pts = doc.add_paragraph("Marking points: " + "; ".join(q["key_points"]))
                    pts.paragraph_format.left_indent = Pt(24)
            number += 1

    if isinstance(path, (str, os.PathLike)):
        folder = os.path.dirname(path)
        if folder:
            os.makedirs(folder, exist_ok=True)
    doc.save(path)  # path can also be an in-memory buffer


def docx_bytes(paper: dict, **kwargs) -> bytes:
    """Return the Word file as bytes, for download buttons in the web app."""
    buffer = io.BytesIO()
    export_docx(paper, buffer, **kwargs)
    return buffer.getvalue()
